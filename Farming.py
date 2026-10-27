"""
Farming.py - Priorizacion de sellers para el equipo de Farming.

Farming potencia a sellers que ya venden en Paris. Cada account manager (AM)
ve una categoria (division). Esta hoja le entrega a cada AM su lista de
sellers potenciables, ordenada por prioridad, con la variable debil a atacar
y el motivo (el dato del seller que la explica).

Como se arma
------------
1. Una fila por seller. Un seller mixto tiene score en duros y en blandos; se
   usa el de la dureza de su categoria principal. Si no tiene score en esa
   dureza, se usa su fila con mayor venta.

2. Filtro "potenciable" (deben cumplirse todos):
     estado_seller         in {Activo, Reactivado, Dormant}   ya vende o se apago hace poco
     segmento              in {Medio - Alto, Medio}           tiene base y espacio para crecer
                                                               (Alto -> Personalizado BD,
                                                                Bajo -> mails automaticos)
     categoria_reputacion  in {Medio, Bueno, Destacado, Plan Growth}
                                                               no potenciar a quien da mala experiencia
     sku_creados           >= 5                                catalogo minimo para trabajar
     n_banderas            >= 1                                tiene al menos una variable debil que atacar

3. Prioridad dentro de cada categoria: score de mayor a menor (desempate: venta).
   prioridad = 1 es el primer seller que deberia contactar el AM de esa categoria.
   Si se define POOL_SEMANAL, solo quedan los N primeros de cada categoria.
"""

import math

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import DUROS, BLANDOS
from Perfil import CORTE_ALTO, CORTE_MEDIO_ALTO, CORTE_MEDIO, COLOR_SEGMENTO
from Tratamiento import filtrar_sp


# ---- parametros de priorizacion (editar aca si cambian) ----
ESTADOS_FARMING = ["Activo", "Reactivado", "Dormant"]
SEGMENTOS_FARMING = ["Medio - Alto", "Medio"]
REPUTACION_FARMING = ["Medio", "Bueno", "Destacado", "Plan Growth"]
SKU_MIN_FARMING = 5
POOL_SEMANAL = None   # ej: 20 -> solo los 20 primeros de cada categoria. None = todos los potenciables

NOMBRE_HOJA = "Priorización de farming"

COLS_FARMING = ["categoria", "prioridad", "seller_id", "name", "admin_email",
                "pct_skus_categoria", "estado_seller", "categoria_reputacion", "division_db",
                "score", "segmento", "pts_para_subir_segmento",
                "variable_debil", "motivo", "accionable", "palanca",
                "flags", "n_banderas", "regalo_sp",
                "venta", "venta_percentil",
                "dias_vendidos", "tramos_con_venta", "percentil_recurrencia",
                "cvr_seller", "percentil_cvr",
                "sku_creados", "sku_stock", "pct_con_stock",
                "percentil_stock", "percentil_volumen_stock",
                "invirtio", "total_invertido"]


def _dureza_de(division) -> str:
    if division in DUROS:
        return "duros"
    if division in BLANDOS:
        return "blandos"
    return "otros"


def _pts_para_subir(fila) -> float:
    """Puntos de score que le faltan para pasar al segmento siguiente."""
    siguiente = {"Medio - Alto": CORTE_ALTO, "Medio": CORTE_MEDIO_ALTO, "Bajo": CORTE_MEDIO}
    corte = siguiente.get(fila["segmento"])
    if corte is None or pd.isna(fila["score"]):
        return None
    return math.ceil((corte - fila["score"]) * 10) / 10   # hacia arriba: 67.95 -> le falta 0.1, no 0.0


def una_fila_por_seller(df: pd.DataFrame) -> pd.DataFrame:
    """Deja una fila por seller: la de la dureza de su categoria principal."""
    df = df.copy()
    df["categoria"] = df["categoria"].fillna("Sin categoria")
    es_su_dureza = df["division_db"] == df["categoria"].apply(_dureza_de)
    df["_orden"] = es_su_dureza.astype(int)
    df = df.sort_values(["seller_id", "_orden", "venta"], ascending=[True, False, False])
    return df.drop_duplicates("seller_id").drop(columns="_orden")


def priorizar_farming(df: pd.DataFrame) -> pd.DataFrame:
    """Sellers potenciables para Farming, con prioridad dentro de su categoria."""
    # marca de Regalo SP (contexto para el AM), antes de dejar una fila por seller
    sp = filtrar_sp(df)
    regla_sp = sp.set_index(["seller_id", "division_db"])["regla_sp"]

    base = una_fila_por_seller(df)

    potenciable = (
        base["estado_seller"].isin(ESTADOS_FARMING)
        & base["segmento"].isin(SEGMENTOS_FARMING)
        & base["categoria_reputacion"].isin(REPUTACION_FARMING)
        & (base["sku_creados"] >= SKU_MIN_FARMING)
        & (base["n_banderas"] >= 1)
    )
    farm = base[potenciable].copy()

    farm["regla_sp"] = [regla_sp.get((s, d), "") for s, d in zip(farm["seller_id"], farm["division_db"])]
    farm["regalo_sp"] = ["No" if r == "" else f"Si ({r})" for r in farm["regla_sp"]]
    farm["pts_para_subir_segmento"] = farm.apply(_pts_para_subir, axis=1)

    farm = farm.sort_values(["categoria", "score", "venta"], ascending=[True, False, False])
    farm["prioridad"] = farm.groupby("categoria").cumcount() + 1
    if POOL_SEMANAL:
        farm = farm[farm["prioridad"] <= POOL_SEMANAL]
    return farm


def escribir_hoja_farming(df: pd.DataFrame, path: str = "score_growth.xlsx") -> pd.DataFrame:
    farm = priorizar_farming(df)
    tabla = farm[[c for c in COLS_FARMING if c in farm.columns]]

    wb = load_workbook(path)
    if NOMBRE_HOJA in wb.sheetnames:
        del wb[NOMBRE_HOJA]
    ws = wb.create_sheet(NOMBRE_HOJA)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    cols = list(tabla.columns)
    for j, col in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = border

    formatos = {"venta": "$#,##0", "total_invertido": "$#,##0", "score": "0",
                "pts_para_subir_segmento": "0.0", "pct_skus_categoria": "0%"}
    texto_largo = {"motivo", "accionable", "palanca", "flags", "name", "admin_email"}

    for i, (_, fila) in enumerate(tabla.iterrows(), start=2):
        for j, col in enumerate(cols, start=1):
            val = fila[col]
            if pd.isna(val):
                val = None
            c = ws.cell(row=i, column=j, value=val)
            c.font = cell_font
            c.alignment = left if col in texto_largo else center
            c.border = border
            if col in formatos and val is not None:
                c.number_format = formatos[col]
        color = COLOR_SEGMENTO.get(fila.get("segmento"))
        if color and "segmento" in cols:
            ws.cell(row=i, column=cols.index("segmento") + 1).fill = PatternFill("solid", fgColor=color)

    anchos = {"motivo": 60, "accionable": 38, "palanca": 32, "flags": 34,
              "name": 28, "admin_email": 30, "categoria": 18, "variable_debil": 20}
    for j, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = anchos.get(col, min(max(len(col), 12) + 2, 26))

    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(tabla) + 1}"
    ws.freeze_panes = "D2"   # deja fijas categoria, prioridad y seller_id
    wb.save(path)

    print(f"Hoja '{NOMBRE_HOJA}' agregada: {len(tabla)} sellers potenciables")
    resumen_farming(farm)
    return farm


def resumen_farming(farm: pd.DataFrame):
    """Conteos en consola: cuantos sellers tiene cada AM y que variable atacar."""
    with pd.option_context("display.max_columns", None, "display.width", 250):
        print("\n=== FARMING: SELLERS POTENCIABLES POR CATEGORIA Y SEGMENTO ===")
        print(pd.crosstab(farm["categoria"], farm["segmento"], margins=True, margins_name="Total"))
        print("\n=== FARMING: VARIABLE DEBIL A ATACAR POR CATEGORIA ===")
        print(pd.crosstab(farm["categoria"], farm["variable_debil"], margins=True, margins_name="Total"))


if __name__ == "__main__":
    # Requiere que score_growth.xlsx tenga la hoja "Score Growth" con categoria, y
    # recalcula el perfilado para tener banderas y motivo.
    from Perfil import perfilar
    RUTA = "score_growth.xlsx"
    df = perfilar(pd.read_excel(RUTA, sheet_name="Score Growth"))
    escribir_hoja_farming(df, RUTA)
    print("\nListo.")
